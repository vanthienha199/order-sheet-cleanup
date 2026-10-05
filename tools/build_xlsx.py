"""Builds local-run/halvorsen_orders.xlsx: the raw export, the cleaned table, a summary and a change log,
styled the way apps-script/Code.gs styles the live sheet (Halvorsen Office Supply ledger look).
Input comes from `node tools/local-run.js`, which runs the same Clean.gs logic locally."""
import csv
import json
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

root = Path(__file__).resolve().parent.parent
run = root / "local-run"
INK, MUTED, LEDGER, YELLOW, LINE = "111111", "5E5C57", "F7F6F1", "F2B705", "E4E2DA"
DISPLAY, BODY, MONO = "Zilla Slab", "IBM Plex Sans", "IBM Plex Mono"

head_fill = PatternFill("solid", fgColor=INK)
head_font = Font(name=DISPLAY, bold=True, color="FFFFFF", size=11)
body_font = Font(name=BODY, size=10, color=INK)
muted_font = Font(name=BODY, size=10, color=MUTED, italic=True)
hair = Border(bottom=Side(style="thin", color=LINE))


def read_tsv(name):
    with open(run / name, newline="") as f:
        return list(csv.reader(f, delimiter="\t"))


def num(v):
    try:
        return float(v) if "." in v else int(v)
    except ValueError:
        return v


wb = Workbook()

# Raw export, left exactly as it arrives, in the default look.
raw = wb.active
raw.title = "Raw orders"
for row in read_tsv("raw.tsv"):
    raw.append(row)
for col, width in zip("ABCDEFGHIJK", [12, 14, 20, 30, 12, 14, 26, 8, 12, 12, 12]):
    raw.column_dimensions[col].width = width

# Clean orders, styled.
clean = wb.create_sheet("Clean orders")
rows = read_tsv("clean.tsv")
clean.append(rows[0])
for r in rows[1:]:
    clean.append([r[0], datetime.strptime(r[1], "%Y-%m-%d").date() if r[1][:2] == "20" else r[1], *r[2:7], num(r[7]), num(r[8]), num(r[9]), r[10], r[11]])
for cell in clean[1]:
    cell.fill, cell.font = head_fill, head_font
    cell.alignment = Alignment(vertical="center")
clean.row_dimensions[1].height = 26
n = len(rows)
for row in clean.iter_rows(min_row=2, max_row=n):
    for cell in row:
        cell.font = body_font
        cell.border = hair
    row[1].number_format = "yyyy-mm-dd"
    row[8].number_format = row[9].number_format = '"$"#,##0.00'
    row[11].font = muted_font
clean.conditional_formatting.add(f"A2:L{n}", FormulaRule(formula=['LEFT($L2,6)="Check:"'], fill=PatternFill("solid", bgColor=YELLOW), font=Font(color=INK)))
clean.freeze_panes = "A2"
for col, width in zip("ABCDEFGHIJKL", [11, 12, 18, 30, 11, 15, 26, 6, 11, 11, 11, 34]):
    clean.column_dimensions[col].width = width

# Summary from the same numbers the weekly email uses.
data = json.loads((run / "summary.json").read_text())
s, stats = data["summary"], data["stats"]
k = s["kpis"]
summ = wb.create_sheet("Summary")
summ.sheet_view.showGridLines = False
summ.column_dimensions["A"].width = 3
for col, width in zip("BCDEFGH", [28, 12, 16, 12, 4, 16, 16]):
    summ.column_dimensions[col].width = width
summ["B2"] = "Halvorsen Office Supply"
summ["B2"].font = Font(name=BODY, size=10, bold=True, color=MUTED)
summ["B3"] = "Week of Sep 28 to Oct 4, 2026"
summ["B3"].font = Font(name=DISPLAY, size=20, bold=True, color=INK)
up = round(k["revenueDelta"] * 100)
summ["B4"] = f"Net revenue ${k['revenue']:,.2f} from {k['orders']} orders, up {up}% on last week."
summ["B4"].font = Font(name=BODY, size=12, color=INK)

def table(top, title, header, body, money_cols):
    summ.cell(row=top, column=2, value=title).font = Font(name=DISPLAY, size=13, bold=True, color=INK)
    for j, h in enumerate(header):
        c = summ.cell(row=top + 1, column=2 + j, value=h)
        c.fill, c.font = head_fill, Font(name=BODY, bold=True, color="FFFFFF", size=10)
    for i, line in enumerate(body):
        for j, v in enumerate(line):
            c = summ.cell(row=top + 2 + i, column=2 + j, value=v)
            c.font, c.border = body_font, hair
            if j in money_cols:
                c.number_format = '"$"#,##0.00'
    return top + 2 + len(body)

end = table(6, "Revenue by week", ["Week starting", "Orders", "Net revenue", "Refunds"], s["weekly"], {2})
chart = BarChart()
chart.title = "Net revenue by week"
chart.style = 1
chart.legend = None
chart.add_data(Reference(summ, min_col=4, min_row=7, max_row=end - 1), titles_from_data=True)
chart.set_categories(Reference(summ, min_col=2, min_row=8, max_row=end - 1))
chart.series[0].graphicalProperties.solidFill = INK
chart.height, chart.width = 7, 14
summ.add_chart(chart, "G6")
end = table(end + 2, "Top products", ["Product", "Units", "Net revenue"], s["products"][:6], {2})
table(end + 2, "Revenue by region", ["Region", "Net revenue"], s["regions"], {1})

# Change log: one honest line for the run that produced this sheet.
log = wb.create_sheet("Change log")
log.append(["When", "What ran", "Raw rows", "Clean orders", "Fields fixed", "Duplicates removed", "Blank rows skipped", "Flagged for a person"])
log.append(["Oct 5, 2026", "Clean.gs cleanup, run locally", stats["input"], stats["output"], stats["fixes"], stats["duplicates"], stats["blank"], stats["flagged"]])
for cell in log[1]:
    cell.fill, cell.font = head_fill, Font(name=DISPLAY, bold=True, color="FFFFFF", size=11)
for cell in log[2]:
    cell.font, cell.border = body_font, hair
    if isinstance(cell.value, int):
        cell.number_format = "#,##0"
log.sheet_view.showGridLines = False
for col, width in zip("ABCDEFGH", [14, 30, 12, 15, 15, 23, 22, 23]):
    log.column_dimensions[col].width = width

out = run / "halvorsen_orders.xlsx"
wb.save(out)
print(out, out.stat().st_size, "bytes")
